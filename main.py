import os
import glob
import shutil
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

def backup_file(file_path):
    """ファイルのバックアップを作成する"""
    new_file_path = os.path.join(os.path.dirname(file_path), "bk" + os.path.basename(file_path))
    shutil.copy(file_path, new_file_path)
    print(f"Backup created: {new_file_path}")

def process_word(file_path):
    """Wordファイルの個人情報を削除する"""
    doc = Document(file_path)
    # Wordファイルのプロパティから個人情報を削除
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(file_path)
    print(f"Processed Word file: {file_path}")

def process_excel(file_path):
    """Excelファイルの個人情報を削除する"""
    wb = load_workbook(file_path)
    # Excelファイルのプロパティから個人情報を削除
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    wb.save(file_path)
    print(f"Processed Excel file: {file_path}")

def process_powerpoint(file_path):
    """PowerPointファイルの個人情報を削除する"""
    prs = Presentation(file_path)
    # PowerPointファイルのプロパティから個人情報を削除
    prs.core_properties.author = ""
    prs.core_properties.last_modified_by = ""
    prs.save(file_path)
    print(f"Processed PowerPoint file: {file_path}")

def main():
    for file_path in glob.glob('./*.*'):
        if file_path.endswith(('.docx', '.doc')):
            backup_file(file_path)
            process_word(file_path)
        elif file_path.endswith(('.xlsx', '.xls')):
            backup_file(file_path)
            process_excel(file_path)
        elif file_path.endswith(('.pptx', '.ppt')):
            backup_file(file_path)
            process_powerpoint(file_path)

if __name__ == "__main__":
    main()
